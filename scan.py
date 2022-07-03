import os
import pdfplumber
from docx import Document
from openpyxl import load_workbook

# supported file types
supported_files = ['.docx', '.pdf', '.xlsx', '.txt']


# Returns all the text in a file as a list
# Return an empty list if no text is found
def get_file_text(file) -> list:
    """

    :rtype: list
    """
    # List object to return any text found
    file_text = []

    # split filename into name and extension
    name, ext = os.path.splitext(file)

    if ext.lower() == '.pdf':
        # this is a pdf file - open it with pdfplumber
        # https://github.com/jsvine/pdfplumber
        with pdfplumber.open(file) as pdf:
            # loop through each page in the pdf
            for p in pdf.pages:
                # extract the text as a str
                text = p.extract_text()
                for line in text.split('\n'):
                    file_text.append(line)
    elif ext.lower() == '.docx':
        # this is a docx file
        # https://python-docx.readthedocs.io/en/latest/user/text.html
        document = Document(file)
        for p in document.paragraphs:
            file_text.append(p.text)
    elif ext.lower() == '.xlsx':
        # This is a xlsx
        # https://openpyxl.readthedocs.io/en/stable/tutorial.html#loading-from-a-file
        wb = load_workbook(file)
        # scan through all the worksheets
        for ws in wb:
            for row in ws.values:
                for value in row:
                    file_text.append(value)
    elif ext.lower() == '.txt':
        # This is a text file
        with open(file) as f:
            for line in f.readlines():
                file_text.append(line.rstrip())
    else:
        raise RuntimeError('Unsupported file type: ' + ext)
    return file_text


def scan_files() -> list:
    # list of files found
    found_list = []

    # Walk all files starting from the files directory
    for root, dirs, files in os.walk('files'):
        # look for supported file types
        for name in files:
            file, ext = os.path.splitext(name)
            # if the file extension is in the supported list
            # add it to the found list.  Handle case sensitivity
            if ext.lower() in supported_files:
                found_list.append(os.path.join(root, name))
    return found_list


def show_aggie_pride():

    slogan_list = ['Aggie Pride - Worldwide',
                    'Aggie Pride!',
                    'And That\'s On My 1891',
                    'Aggies ❤️ engineering',
                   'Aggies Do!',
                   'I like my Aggie Pride!'
                  ]

    slogan_list = ['Aggie Pride - Worldwide', 'Aggies Do!', 'Aggie Pride!', 'And That\'s On My 1891',
                   'Aggies ❤️ engineering', 'Aggies Do!', 'Aggies Get it Done!', 'NCAT',
                   'Aggie born aggie bred, when I die Ill be aggie dead', 'Aggies Do']

    return slogan_list


if __name__ == '__main__':
    print(show_aggie_pride())
    print('---')
    print(scan_files())