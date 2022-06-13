import unittest
import pdfplumber
import os
import re
from scan import show_aggie_pride, scan_files
import spacy
from docx import Document
from openpyxl import load_workbook


class ScanTests(unittest.TestCase):
    def test_aggie_pride(self):
        # get the slogans
        slogans = show_aggie_pride()

        self.assertIn('Aggie Pride - Worldwide', slogans)

    def test_pdf_parse(self):
        # get a platform safe path to a test pdf file
        pdf_file = os.path.join('files', 'november_statement.pdf')

        # open the file
        with pdfplumber.open(pdf_file) as pdf:
            # loop through each page in the pdf
            for p in pdf.pages:
                # extract the text as a str
                text = p.extract_text()

                # Check for the name that should be in this PDF
                self.assertIn('Name: John Smith', text)

                # Check for the account number
                self.assertIn('Account Number: 3942-29992', text)

                # Regex match on the account number
                m = re.search(r'\d+-\d+', text)
                self.assertTrue(m)

    def test_ssn(self):
        # Walk all files starting from the files directory
        for root, dirs, files in os.walk('files'):
            for name in files:
                # Look for files that should have ssn in them
                if name == 'ss_info.pdf':
                    # open the file
                    with pdfplumber.open(os.path.join(root, name)) as pdf:
                        # loop through each page in the pdf
                        for p in pdf.pages:
                            # extract the text as a str
                            text = p.extract_text()
                            # check for the ssn
                            m = re.search(r'\d{3}-\d{2}-\d{4}', text)
                            self.assertTrue(m)

    def test_scan_files(self):
        # Test to make sure scan_files returns the expected results
        expected_result = ['files/november_statement.pdf',
                           'files/Documents/twitter_info.docx',
                           'files/Documents/Statements/Retirement/ss_info.pdf',
                           'files/Downloads/address_book.xlsx',
                           'files/Downloads/address_book.txt']

        # Make expected_result os safe by checking the seperator
        if os.sep != '/':
            for i in range(len(expected_result)):
                expected_result[i] = expected_result[i].replace('/', os.sep)

        # Make sure all the expected files are actually found
        scan = scan_files()
        for f in expected_result:
            self.assertIn(f, scan)

    def test_name_recognition(self):
        # python -m spacy download en_core_web_sm
        try:
            nlp = spacy.load("en_core_web_sm")
        except OSError:
            from spacy.cli import download
            download("en_core_web_sm")
            nlp = spacy.load("en_core_web_sm")

        # Create a sample string with a mix of things that look like they could be names
        names_string = 'Name: John Jones\nAddress: 123 John W Mitchell Drive\nNorth Carolina\nHarsh Mupali'
        doc = nlp(names_string)

        # expected results
        # Be careful with addresses that look like a name - they are often detected incorrectly
        expected = {'John Jones': 'PERSON',
                    '123': 'CARDINAL', 'John W Mitchell Drive': 'PERSON',
                    'North Carolina': 'GPE',
                    'Harsh Mupali': 'PERSON'}

        for ent in doc.ents:
            self.assertEqual(expected[ent.text], ent.label_)

    def test_docx(self):
        # test to make sure we can read a docx ok
        doc = 'files/Documents/twitter_info.docx'

        # Fix seperator for windows (or other platforms)
        if os.sep != '/':
            doc = doc.replace('/', os.sep)

        # Read test document which contains a twitter handle
        document = Document(doc)
        for p in document.paragraphs:
            m = re.search(r'(@\w+)', p.text)
            self.assertEqual('@john_jones', m.group(1))

    def test_xlsx(self):
        xlsx = 'files/Downloads/address_book.xlsx'

        # Fix seperator for windows (or other platforms)
        if os.sep != '/':
            xlsx = xlsx.replace('/', os.sep)

        # load a workbook
        wb = load_workbook(xlsx)
        # scan through all the worksheets
        phones = []
        for ws in wb:
            for row in ws.values:
                for value in row:
                    # Find phone numbers
                    m = re.search(r'(\d{3}-\d{3}-\d{4})', value)
                    if m:
                        phones.append(m.group(1))

        # check to make sure we found all the phone numbers
        self.assertIn('336-555-1212', phones)
        self.assertIn('919-555-1212', phones)
        self.assertIn('970-555-1212', phones)

    def test_txt(self):
        txt = 'files/Downloads/address_book.txt'

        # Fix seperator for windows (or other platforms)
        if os.sep != '/':
            txt = txt.replace('/', os.sep)

        names = []
        with open(txt) as f:
            for line in f.readlines():
                m = re.search(r'([A-Z][a-z]+ [A-Z][a-z]+)', line)
                if m:
                    names.append(m.group(1))

        # check to make sure we found all the names
        self.assertIn('John Jones', names)
        self.assertIn('James Johnson', names)
        self.assertIn('Roger Jones', names)


if __name__ == '__main__':
    unittest.main()
